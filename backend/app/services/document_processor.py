import os
import logging
from typing import Dict, Any, Tuple, Optional
import fitz  # PyMuPDF
import docx
from app.services.ocr_service import extract_text_from_image, is_ocr_available

logger = logging.getLogger(__name__)


def process_uploaded_document(
    file_path: str,
    original_filename: str,
    mime_type: str
) -> Tuple[str, Dict[str, Any], str, Optional[str]]:
    """
    Extracts text and metadata from uploaded file.
    Returns: (extracted_text, metadata, status, error_message)
    status can be: 'completed', 'completed_with_warning', or 'failed'
    """
    ext = os.path.splitext(original_filename)[1].lower()
    extracted_text = ""
    metadata: Dict[str, Any] = {
        "file_extension": ext,
        "original_filename": original_filename,
    }
    error_message = None
    status = "completed"

    try:
        # 1. PDF Documents
        if ext == ".pdf" or mime_type == "application/pdf":
            doc = fitz.open(file_path)
            metadata["page_count"] = len(doc)
            text_parts = []
            for page_num in range(len(doc)):
                page = doc[page_num]
                text = page.get_text()
                if text.strip():
                    text_parts.append(f"--- Page {page_num + 1} ---\n{text.strip()}")
            doc.close()

            extracted_text = "\n\n".join(text_parts)
            # If PDF is scanned (has very little text), mention OCR needed
            if len(extracted_text.strip()) < 50 and metadata["page_count"] > 0:
                metadata["is_scanned"] = True
                status = "completed_with_warning"
                error_message = "PDF appears to be a scanned document with minimal embedded text."

        # 2. Word Documents (.docx)
        elif ext == ".docx" or "wordprocessingml" in mime_type:
            doc = docx.Document(file_path)
            text_parts = []
            for p in doc.paragraphs:
                if p.text.strip():
                    text_parts.append(p.text.strip())
            for t in doc.tables:
                table_rows = []
                for row in t.rows:
                    row_cells = [cell.text.strip() for cell in row.cells]
                    table_rows.append(" | ".join(row_cells))
                if table_rows:
                    text_parts.append("\n[Table Data]\n" + "\n".join(table_rows))
            extracted_text = "\n\n".join(text_parts)
            metadata["paragraph_count"] = len(doc.paragraphs)
            metadata["table_count"] = len(doc.tables)

        # 3. Images (OCR)
        elif ext in [".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tiff"] or mime_type.startswith("image/"):
            text, success, err = extract_text_from_image(file_path)
            extracted_text = text
            metadata["ocr_applied"] = True
            if not success:
                status = "completed_with_warning"
                error_message = err
            else:
                metadata["ocr_character_count"] = len(text)

        # 4. Text & Code files (.txt, .md, .py, .java, .cpp, .c, .js, .ts, .html, .css, .json, .sql)
        else:
            # Try utf-8 first, fallback to latin-1
            try:
                with open(file_path, "r", encoding="utf-8") as f:
                    extracted_text = f.read()
            except UnicodeDecodeError:
                with open(file_path, "r", encoding="latin-1") as f:
                    extracted_text = f.read()

            metadata["line_count"] = len(extracted_text.splitlines())
            metadata["is_code"] = ext in [".py", ".java", ".c", ".cpp", ".js", ".ts", ".html", ".css", ".json", ".sql"]

        metadata["word_count"] = len(extracted_text.split())
        metadata["char_count"] = len(extracted_text)

        if not extracted_text.strip() and not error_message:
            status = "completed_with_warning"
            error_message = "No readable text could be extracted from this file."

    except Exception as e:
        status = "failed"
        error_message = f"File processing failed: {str(e)}"
        logger.exception(error_message)

    return extracted_text, metadata, status, error_message
