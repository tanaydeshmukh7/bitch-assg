import io
import os
import re
import html
from typing import Dict, Any, List
from datetime import datetime

# DOCX generation
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

# PDF generation using ReportLab
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import inch, cm
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable, Preformatted
)
from reportlab.pdfgen import canvas


class NumberedCanvas(canvas.Canvas):
    """Adds running headers and 'Page X of Y' footers to ReportLab PDF."""
    def __init__(self, *args, **kwargs):
        super(NumberedCanvas, self).__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super(NumberedCanvas, self).showPage()
        super(NumberedCanvas, self).save()

    def draw_page_decorations(self, page_count):
        # Don't draw headers/footers on cover page (page 1) if document has multiple pages
        if self._pageNumber == 1 and page_count > 1:
            return

        self.saveState()
        self.setFont("Helvetica", 9)
        self.setFillColor(colors.HexColor("#64748B"))

        # Header rule and title
        self.drawString(54, 842 - 36, "Assignment Studio")
        self.setStrokeColor(colors.HexColor("#E2E8F0"))
        self.setLineWidth(0.5)
        self.line(54, 842 - 42, 595 - 54, 842 - 42)

        # Footer page number
        page_str = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(595 - 54, 36, page_str)
        self.restoreState()


def strip_html_tags(text: str) -> str:
    """Removes HTML tags and unescapes HTML entities."""
    clean = re.sub(r'<[^>]+>', '', text)
    return html.unescape(clean).strip()


def parse_html_to_blocks(html_content: str) -> List[Dict[str, Any]]:
    """
    Parses section HTML content into canonical blocks:
    - heading (level, text)
    - paragraph (text)
    - list (type: bullet|number, items)
    - code (code_string)
    - table (rows: list of lists)
    """
    blocks = []
    if not html_content:
        return blocks

    # Replace <br> with newlines
    content = re.sub(r'<br\s*/?>', '\n', html_content)

    # 1. Extract pre/code blocks
    code_pattern = re.compile(r'<pre(?:[^>]*)><code(?:[^>]*)>(.*?)</code></pre>', re.DOTALL | re.IGNORECASE)
    parts = []
    last_end = 0
    for match in code_pattern.finditer(content):
        start, end = match.span()
        if start > last_end:
            parts.append(("html", content[last_end:start]))
        parts.append(("code", match.group(1)))
        last_end = end
    if last_end < len(content):
        parts.append(("html", content[last_end:]))

    for part_type, part_data in parts:
        if part_type == "code":
            clean_code = html.unescape(part_data)
            # Remove any residual inner tags
            clean_code = re.sub(r'<[^>]+>', '', clean_code)
            blocks.append({"type": "code", "content": clean_code})
            continue

        # Process HTML sub-parts
        # Check for headings
        chunk = part_data
        tag_pattern = re.compile(r'<(h[1-6]|p|ul|ol|table)(?:[^>]*)>(.*?)</\1>', re.DOTALL | re.IGNORECASE)
        last_idx = 0
        found_tags = False

        for tag_match in tag_pattern.finditer(chunk):
            found_tags = True
            tag_name = tag_match.group(1).lower()
            inner = tag_match.group(2)

            if tag_name.startswith('h'):
                level = int(tag_name[1])
                blocks.append({"type": "heading", "level": level, "text": strip_html_tags(inner)})
            elif tag_name == 'p':
                p_text = strip_html_tags(inner)
                if p_text:
                    blocks.append({"type": "paragraph", "text": p_text})
            elif tag_name in ['ul', 'ol']:
                list_type = "bullet" if tag_name == 'ul' else "number"
                items = re.findall(r'<li(?:[^>]*)>(.*?)</li>', inner, re.DOTALL | re.IGNORECASE)
                clean_items = [strip_html_tags(it) for it in items if strip_html_tags(it)]
                if clean_items:
                    blocks.append({"type": "list", "list_type": list_type, "items": clean_items})
            elif tag_name == 'table':
                rows = []
                for tr in re.findall(r'<tr(?:[^>]*)>(.*?)</tr>', inner, re.DOTALL | re.IGNORECASE):
                    cells = re.findall(r'<(?:td|th)(?:[^>]*)>(.*?)</(?:td|th)>', tr, re.DOTALL | re.IGNORECASE)
                    rows.append([strip_html_tags(c) for c in cells])
                if rows:
                    blocks.append({"type": "table", "rows": rows})

        # Fallback if no specific tags were matched
        if not found_tags and strip_html_tags(chunk):
            for para in chunk.split('\n\n'):
                txt = strip_html_tags(para)
                if txt:
                    blocks.append({"type": "paragraph", "text": txt})

    return blocks


def extract_content_text(content_obj: Any) -> str:
    """Extracts raw HTML string from content object (string or dict)."""
    if isinstance(content_obj, str):
        return content_obj
    elif isinstance(content_obj, dict):
        if "html" in content_obj:
            return content_obj["html"]
        elif "text" in content_obj:
            return content_obj["text"]
    return ""


# ====================================================================
# DOCX EXPORT GENERATOR
# ====================================================================

def generate_docx_document(
    project_title: str,
    subject: str,
    metadata: Dict[str, Any],
    sections: List[Dict[str, Any]],
    include_cover_page: bool = True
) -> bytes:
    """Generates an editable Word (.docx) document using native styles."""
    doc = docx.Document()

    # Set page to A4 with 1-inch margins
    for section in doc.sections:
        section.page_width = Inches(8.27)   # 210mm
        section.page_height = Inches(11.69) # 297mm
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Primary colors
    PRIMARY_COLOR = RGBColor(37, 99, 235)  # #2563EB
    DARK_TEXT = RGBColor(15, 23, 42)       # #0F172A
    MUTED_TEXT = RGBColor(100, 116, 139)   # #64748B

    college = metadata.get("college")
    department = metadata.get("department")
    class_name = metadata.get("class_name") or metadata.get("class")
    student_name = metadata.get("student_name")
    roll_number = metadata.get("roll_number")
    teacher_name = metadata.get("teacher_name")

    # --- 1. COVER PAGE ---
    if include_cover_page:
        if college:
            p_col = doc.add_paragraph()
            p_col.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_col = p_col.add_run(college.upper())
            run_col.bold = True
            run_col.font.size = Pt(16)
            run_col.font.color.rgb = DARK_TEXT

        if department:
            p_dept = doc.add_paragraph()
            p_dept.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_dept = p_dept.add_run(department)
            run_dept.font.size = Pt(12)
            run_dept.font.color.rgb = MUTED_TEXT

        doc.add_paragraph()  # Spacing

        # Title
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_title = p_title.add_run(project_title)
        run_title.bold = True
        run_title.font.size = Pt(24)
        run_title.font.color.rgb = PRIMARY_COLOR

        # Subject
        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_sub = p_sub.add_run(f"Subject: {subject}")
        run_sub.font.size = Pt(14)
        run_sub.font.color.rgb = DARK_TEXT

        doc.add_paragraph()
        doc.add_paragraph()

        # Metadata Table (Submitted by / Submitted to)
        meta_table = doc.add_table(rows=2, cols=2)
        meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        meta_table.autofit = True

        cell_left = meta_table.cell(0, 0)
        p_left_header = cell_left.paragraphs[0]
        run_lh = p_left_header.add_run("SUBMITTED BY:")
        run_lh.bold = True
        run_lh.font.size = Pt(10)
        run_lh.font.color.rgb = PRIMARY_COLOR

        p_left_body = cell_left.add_paragraph()
        p_left_body.add_run(f"Name: {student_name or '________________'}\n")
        p_left_body.add_run(f"Roll No: {roll_number or '________________'}\n")
        if class_name:
            p_left_body.add_run(f"Class: {class_name}")

        cell_right = meta_table.cell(0, 1)
        p_right_header = cell_right.paragraphs[0]
        run_rh = p_right_header.add_run("SUBMITTED TO:")
        run_rh.bold = True
        run_rh.font.size = Pt(10)
        run_rh.font.color.rgb = PRIMARY_COLOR

        p_right_body = cell_right.add_paragraph()
        p_right_body.add_run(f"Instructor: {teacher_name or '________________'}\n")
        p_right_body.add_run(f"Date: {datetime.now().strftime('%B %d, %Y')}")

        doc.add_page_break()

    # --- 2. DOCUMENT SECTIONS ---
    for sec_idx, sec in enumerate(sections):
        title = sec.get("title", f"Section {sec_idx + 1}")
        content_html = extract_content_text(sec.get("content", ""))

        # Section Heading
        h = doc.add_heading(level=1)
        run_h = h.add_run(title)
        run_h.font.color.rgb = PRIMARY_COLOR
        run_h.font.size = Pt(16)
        run_h.bold = True

        blocks = parse_html_to_blocks(content_html)
        if not blocks:
            # Fallback if plain text without HTML
            p = doc.add_paragraph(strip_html_tags(content_html) or "No content recorded for this section.")
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(6)

        for block in blocks:
            b_type = block.get("type")

            if b_type == "heading":
                level = min(block.get("level", 2), 4)
                sub_h = doc.add_heading(level=level)
                sub_h_run = sub_h.add_run(block.get("text", ""))
                sub_h_run.font.color.rgb = DARK_TEXT
                sub_h_run.font.size = Pt(14 if level == 2 else 12)
                sub_h_run.bold = True

            elif b_type == "paragraph":
                p = doc.add_paragraph()
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(6)
                p.add_run(block.get("text", ""))

            elif b_type == "list":
                list_type = block.get("list_type", "bullet")
                style_name = 'List Bullet' if list_type == 'bullet' else 'List Number'
                for item_text in block.get("items", []):
                    try:
                        p_item = doc.add_paragraph(item_text, style=style_name)
                    except Exception:
                        p_item = doc.add_paragraph(f"• {item_text}")
                    p_item.paragraph_format.space_after = Pt(3)

            elif b_type == "code":
                code_text = block.get("content", "")
                table = doc.add_table(rows=1, cols=1)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = table.cell(0, 0)
                # Set light grey background for code block
                shading_elm = parse_xml(r'<w:shd {} w:fill="F1F5F9"/>'.format(nsdecls('w')))
                cell._tc.get_or_add_tcPr().append(shading_elm)

                p_code = cell.paragraphs[0]
                p_code.paragraph_format.space_after = Pt(0)
                run_code = p_code.add_run(code_text)
                run_code.font.name = "Consolas"
                run_code.font.size = Pt(9.5)
                run_code.font.color.rgb = RGBColor(15, 23, 42)

                doc.add_paragraph()  # Spacing after code block

            elif b_type == "table":
                rows = block.get("rows", [])
                if rows:
                    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
                    t.alignment = WD_TABLE_ALIGNMENT.CENTER
                    for r_idx, row in enumerate(rows):
                        for c_idx, val in enumerate(row):
                            cell = t.cell(r_idx, c_idx)
                            cell.text = str(val)
                            if r_idx == 0:
                                # Header styling
                                shd = parse_xml(r'<w:shd {} w:fill="E2E8F0"/>'.format(nsdecls('w')))
                                cell._tc.get_or_add_tcPr().append(shd)
                                for p in cell.paragraphs:
                                    for r in p.runs:
                                        r.bold = True
                    doc.add_paragraph()

        # Small spacing between sections
        doc.add_paragraph()

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


# ====================================================================
# PDF EXPORT GENERATOR (ReportLab A4)
# ====================================================================

def generate_pdf_document(
    project_title: str,
    subject: str,
    metadata: Dict[str, Any],
    sections: List[Dict[str, Any]],
    include_cover_page: bool = True
) -> bytes:
    """Generates an A4 PDF document matching the canonical assignment model."""
    buffer = io.BytesIO()

    # 1. Page template setup: A4 size with 0.75 in margins
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=54,
        rightMargin=54,
        topMargin=54,
        bottomMargin=54
    )

    styles = getSampleStyleSheet()

    # Custom styles
    primary_blue = colors.HexColor("#2563EB")
    dark_slate = colors.HexColor("#0F172A")
    muted_slate = colors.HexColor("#64748B")
    border_color = colors.HexColor("#E2E8F0")

    style_cover_institution = ParagraphStyle(
        'CoverInstitution',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        alignment=1,  # Center
        textColor=dark_slate,
        spaceAfter=6
    )

    style_cover_dept = ParagraphStyle(
        'CoverDept',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=12,
        leading=16,
        alignment=1,
        textColor=muted_slate,
        spaceAfter=40
    )

    style_cover_title = ParagraphStyle(
        'CoverTitle',
        parent=styles['Title'],
        fontName='Helvetica-Bold',
        fontSize=24,
        leading=30,
        alignment=1,
        textColor=primary_blue,
        spaceAfter=15
    )

    style_cover_subject = ParagraphStyle(
        'CoverSubject',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=14,
        leading=18,
        alignment=1,
        textColor=dark_slate,
        spaceAfter=60
    )

    style_section_title = ParagraphStyle(
        'SectionTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=16,
        leading=20,
        textColor=primary_blue,
        spaceBefore=14,
        spaceAfter=8,
        keepWithNext=True
    )

    style_subheading = ParagraphStyle(
        'SubHeading',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=16,
        textColor=dark_slate,
        spaceBefore=10,
        spaceAfter=4,
        keepWithNext=True
    )

    style_body = ParagraphStyle(
        'AcademicBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=15,
        textColor=dark_slate,
        spaceAfter=6
    )

    style_list_item = ParagraphStyle(
        'AcademicList',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=15,
        textColor=dark_slate,
        leftIndent=20,
        spaceAfter=3
    )

    style_code = ParagraphStyle(
        'CodeStyle',
        parent=styles['Code'],
        fontName='Courier',
        fontSize=8.5,
        leading=12,
        textColor=dark_slate
    )

    story = []

    college = metadata.get("college")
    department = metadata.get("department")
    class_name = metadata.get("class_name") or metadata.get("class")
    student_name = metadata.get("student_name")
    roll_number = metadata.get("roll_number")
    teacher_name = metadata.get("teacher_name")

    # --- 1. COVER PAGE ---
    if include_cover_page:
        story.append(Spacer(1, 30))
        if college:
            story.append(Paragraph(college.upper(), style_cover_institution))
        if department:
            story.append(Paragraph(department, style_cover_dept))
        else:
            story.append(Spacer(1, 30))

        story.append(Paragraph(project_title, style_cover_title))
        story.append(Paragraph(f"Subject: {subject}", style_cover_subject))

        # Metadata box (Submitted By / Submitted To)
        meta_data = [
            [
                Paragraph("<b>SUBMITTED BY:</b>", ParagraphStyle('H1', parent=style_body, textColor=primary_blue)),
                Paragraph("<b>SUBMITTED TO:</b>", ParagraphStyle('H2', parent=style_body, textColor=primary_blue))
            ],
            [
                Paragraph(f"<b>Name:</b> {student_name or 'N/A'}<br/><b>Roll No:</b> {roll_number or 'N/A'}<br/><b>Class:</b> {class_name or 'N/A'}", style_body),
                Paragraph(f"<b>Faculty:</b> {teacher_name or 'N/A'}<br/><b>Date:</b> {datetime.now().strftime('%B %d, %Y')}", style_body)
            ]
        ]
        meta_table = Table(meta_data, colWidths=[240, 240])
        meta_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
            ('BOX', (0, 0), (-1, -1), 1, border_color),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, border_color),
            ('TOPPADDING', (0, 0), (-1, -1), 10),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 10),
            ('LEFTPADDING', (0, 0), (-1, -1), 12),
            ('RIGHTPADDING', (0, 0), (-1, -1), 12),
        ]))
        story.append(meta_table)
        story.append(PageBreak())

    # --- 2. SECTIONS ---
    for sec_idx, sec in enumerate(sections):
        title = sec.get("title", f"Section {sec_idx + 1}")
        content_html = extract_content_text(sec.get("content", ""))

        story.append(Paragraph(title, style_section_title))
        story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#E2E8F0"), spaceAfter=8))

        blocks = parse_html_to_blocks(content_html)
        if not blocks:
            raw = strip_html_tags(content_html) or "No section content available."
            story.append(Paragraph(raw, style_body))

        for block in blocks:
            b_type = block.get("type")

            if b_type == "heading":
                story.append(Paragraph(block.get("text", ""), style_subheading))

            elif b_type == "paragraph":
                # Escape XML entities for ReportLab Paragraph
                safe_text = html.escape(block.get("text", "")).replace('&amp;', '&')
                story.append(Paragraph(safe_text, style_body))

            elif b_type == "list":
                is_num = block.get("list_type") == "number"
                for i_idx, itm in enumerate(block.get("items", [])):
                    prefix = f"{i_idx+1}. " if is_num else "• "
                    safe_itm = html.escape(itm).replace('&amp;', '&')
                    story.append(Paragraph(f"{prefix}{safe_itm}", style_list_item))

            elif b_type == "code":
                code_text = block.get("content", "")
                p_code = Preformatted(code_text, style_code)
                code_table = Table([[p_code]], colWidths=[480])
                code_table.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#F1F5F9")),
                    ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                    ('TOPPADDING', (0, 0), (-1, -1), 8),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
                    ('LEFTPADDING', (0, 0), (-1, -1), 10),
                    ('RIGHTPADDING', (0, 0), (-1, -1), 10),
                ]))
                story.append(Spacer(1, 4))
                story.append(code_table)
                story.append(Spacer(1, 6))

            elif b_type == "table":
                rows = block.get("rows", [])
                if rows:
                    tbl_data = []
                    col_count = len(rows[0])
                    col_w = 480 / max(col_count, 1)
                    for r_idx, row in enumerate(rows):
                        row_cells = []
                        for cell_val in row:
                            st = style_body if r_idx > 0 else ParagraphStyle('TH', parent=style_body, fontName='Helvetica-Bold')
                            row_cells.append(Paragraph(html.escape(str(cell_val)), st))
                        tbl_data.append(row_cells)

                    t_flow = Table(tbl_data, colWidths=[col_w] * col_count)
                    t_flow.setStyle(TableStyle([
                        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#F8FAFC")),
                        ('BOX', (0, 0), (-1, -1), 0.5, border_color),
                        ('INNERGRID', (0, 0), (-1, -1), 0.5, border_color),
                        ('TOPPADDING', (0, 0), (-1, -1), 6),
                        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
                    ]))
                    story.append(Spacer(1, 4))
                    story.append(t_flow)
                    story.append(Spacer(1, 6))

        story.append(Spacer(1, 10))

    # Build document with running headers and page numbers
    doc.build(story, canvasmaker=NumberedCanvas)
    buffer.seek(0)
    return buffer.getvalue()
